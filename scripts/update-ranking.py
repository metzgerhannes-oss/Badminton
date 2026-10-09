#!/usr/bin/env python3
"""Weekly DBV *published Excel* ranking snapshots. No tournament scraping.

An age-class rank is DERIVED from the official AKL2 field (U11 contains multiple birth years),
and deliberately not confused with the DBV overall Rank column.
"""
from __future__ import annotations
import io
import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "data" / "ranking.json"
OFFICIAL = "https://turniere.badminton.de"
HISTORY = OFFICIAL + "/ranking/history"
RELEASES = OFFICIAL + "/uploads/ranking/Ranking_{year}_KW{week:02d}.xlsx"
VALID_DISCIPLINES = ("HE", "HD", "HM", "DE", "DD", "DM")
HEADERS = {"dis":"discipline", "rang":"overall_rank", "spielerid":"player_id",
           "gjahr":"birth_year", "akl1":"age_detail","akl2":"age_class",
           "punkte":"points","turniere":"tournaments","gs":"gender",
           "nachname":"last_name","vorname":"first_name"}
MAX_BYTES = 48 * 1024 * 1024
BW_ASSOCIATION = "BAW-Baden-Württemberg"

def iso_week(d: datetime) -> tuple[int,int]:
    i = d.isocalendar()
    return (i.year, i.week)

def clean(value):
    return re.sub(r"\s+", "", str(value or "").strip().lower())

def points(value):
    if isinstance(value, (float,int)):
        return float(value)
    s = str(value or "").strip().replace("\u00a0", "").replace(" ", "")
    if not s:return None
    # DBV German display: 2.670 means 2670, not 2.67.
    if "," in s:
        s=s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+", s):
        s=s.replace(".", "")
    try:return float(s)
    except ValueError:return None

def parse_excel(contents: bytes) -> list[dict]:
    wb=load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    found=[]
    for ws in wb.worksheets:
        iterator=ws.iter_rows(values_only=True)
        header=None
        for index,row in enumerate(iterator):
            fields={clean(cell):j for j,cell in enumerate(row) if cell is not None}
            if "ranglistenplatz" in fields and "rang" not in fields: fields["rang"]=fields["ranglistenplatz"]
            if "points" in fields and "punkte" not in fields: fields["punkte"]=fields["points"]
            if all(x in fields for x in ("dis","spielerid","gjahr","punkte")):
                header=fields
                break
            if index >= 45:break
        if header is None:
            if os.environ.get("DEBUG_RANKING_SCHEMA")=="1":
                samples=[]
                for n,sample in enumerate(ws.iter_rows(values_only=True)):
                    if any(x is not None for x in sample):
                        samples.append([str(x)[:65] if x is not None else "" for x in sample[:18]])
                    if len(samples)>=8 or n>=35:break
                print("Ranking schema mismatch:",ws.title,"samples:",repr(samples),file=sys.stderr)
            continue
        for row in iterator:
            def val(key):
                idx=header.get(key)
                return row[idx] if idx is not None and idx<len(row) else None
            disc=str(val("dis") or "").strip().upper()
            pid=str(val("spielerid") or "").strip()
            p=points(val("punkte"))
            try:year=int(val("gjahr"))
            except (TypeError,ValueError):continue
            if disc not in VALID_DISCIPLINES or not re.fullmatch(r"\d\d-\d{6}",pid) or p is None:
                continue
            overall=val("rang")
            try:overall_rank=int(float(overall))
            except (TypeError,ValueError):overall_rank=None
            found.append({"id":pid,"discipline":disc,"birthYear":year,"points":p,
                "overallRank":overall_rank,"ageClass":str(val("akl2") or ""),
                "ageDetail":str(val("akl1") or ""),
                "gender":str(val("gs") or "").strip().upper(),
                "association":str(val("lvname") or "").strip(),
                "tournaments":int(val("turniere") or 0),
                "firstName":str(val("vorname") or "").strip(),
                "lastName":str(val("nachname") or "").strip()})
    wb.close()
    if not found:raise ValueError("No usable ranking table in official Excel")
    return found

def cohort_ranks(rows: list[dict]) -> dict[tuple[str,str],dict]:
    """Rank by official AKL2 age class (U11 includes 2016, 2017, etc.)."""
    groups=defaultdict(dict)
    for row in rows:
        age_class=str(row.get("ageClass") or "").strip().upper()
        if not re.fullmatch(r"U(?:11|13|15|17|19|22)",age_class):
            continue
        key=(row["discipline"],row["gender"],age_class)
        existing=groups[key].get(row["id"])
        if existing is None or row["points"]>existing["points"]:
            groups[key][row["id"]]=row
    ranked={}
    for (_,_,age_class),members in groups.items():
        order=sorted(members.values(),key=lambda r:-r["points"])
        first_at_points={}
        for index,item in enumerate(order,1):
            first_at_points.setdefault(item["points"],index)
            record=dict(item)
            record["ageClassRank"]=first_at_points[item["points"]]
            record["ageClassSize"]=len(order)
            record["ageClass"]=age_class
            k=(item["id"],item["discipline"])
            if k not in ranked or ranked[k]["ageClassRank"]>record["ageClassRank"]:
                ranked[k]=record
    return ranked

def current_public_export():
    """Read the actual *current* official DBV Excel download (not only old week archives).

    Source HTML provides the authoritative publication KW and timestamp, while
    /ranking/download provides the corresponding complete official XLSX.
    """
    page_url=OFFICIAL+"/ranking"
    export_url=OFFICIAL+"/ranking/download"
    try:
        req=Request(page_url,headers={"User-Agent":"Schmetterlinge-Ranking/1.1 (weekly public export)","Accept":"text/html"})
        with urlopen(req,timeout=25) as resp:
            raw=resp.read(5_000_001)
        if len(raw)>5_000_000:raise ValueError("Ranking page exceeds maximum size")
        html=raw.decode("utf-8","replace")
        # Exact navigation label from the official page, e.g. "Rangliste KW 41".
        match=re.search(r"Rangliste\s+KW\s*(\d{1,2})\b",html,re.I)
        # Published timestamp is needed to determine the correct ISO week year.
        updated_anchor=html.find("zuletzt aktualisiert:")
        updated_match=re.search(r"(\d{2}\.\d{2}\.\d{4})\s+(\d{2}:\d{2}:\d{2})",
            html[updated_anchor:updated_anchor+14000]) if updated_anchor>=0 else None
        if not match or not updated_match:
            raise ValueError("Current ranking week or publication timestamp missing; do not guess")
        source_dt=datetime.strptime(updated_match.group(1)+" "+updated_match.group(2),
            "%d.%m.%Y %H:%M:%S").replace(tzinfo=ZoneInfo("Europe/Berlin"))
        iso_year,published_week=iso_week(source_dt)
        listed_week=int(match.group(1))
        if published_week!=listed_week:
            raise ValueError(f"Ranking label KW {listed_week} conflicts with published date KW {published_week}")
        headers={"User-Agent":"Schmetterlinge-Ranking/1.1 (weekly public export)",
                 "Accept":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream;q=0.9"}
        with urlopen(Request(export_url,headers=headers),timeout=50) as resp:
            if resp.status!=200:raise ValueError("Current export request did not succeed")
            binary=resp.read(MAX_BYTES+1)
            if len(binary)>MAX_BYTES:raise ValueError("Current ranking Excel exceeds size limit")
            if not binary.startswith(b"PK"):raise ValueError("Current ranking download is not XLSX")
        rows=parse_excel(binary)
        print(f"CURRENT DBV RANKING: KW {published_week}/{iso_year}, published {source_dt.isoformat()}, {len(rows)} rows")
        return (iso_year,published_week,export_url,rows,source_dt.isoformat())
    except HTTPError as e:
        if e.code in (401,403,429):
            print(f"Current official DBV export denies automated access: HTTP {e.code}; fall back to archives",file=sys.stderr)
        else:
            print(f"Current DBV export unavailable: HTTP {e.code}; fall back to archives",file=sys.stderr)
    except (URLError,ValueError,OSError,TimeoutError) as e:
        print(f"Current DBV export unavailable: {type(e).__name__}: {e}; fall back to archives",file=sys.stderr)
    return None

def published_export(when:datetime, max_lookback=6):
    for ago in range(max_lookback):
        year,week=iso_week(when-timedelta(weeks=ago))
        url=RELEASES.format(year=year,week=week)
        try:
            req=Request(url,headers={"User-Agent":"Schmetterlinge-Ranking/1.0 (weekly family cohort tracker)","Accept":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
            with urlopen(req,timeout=22) as resp:
                if resp.status!=200:continue
                sz=int(resp.headers.get("Content-Length") or 0)
                if sz>MAX_BYTES:raise ValueError("Excel archive exceeds size limit")
                binary=resp.read(MAX_BYTES+1)
                if len(binary)>MAX_BYTES:raise ValueError("Excel archive exceeds size limit")
            if not binary.startswith(b"PK"):raise ValueError("Not a valid xlsx ZIP file")
            yield year,week,url,parse_excel(binary)
        except HTTPError as e:
            if e.code in (404,410):continue
            if e.code in (401,403,429):
                print(f"Official archive denied automated access: HTTP {e.code}; skipping instead of retrying.",file=sys.stderr)
                return
            print(f"HTTP {e.code} from official archive",file=sys.stderr)
        except (URLError,ValueError,OSError,TimeoutError) as e:
            print(f"Archive unavailable: {type(e).__name__}: {e}",file=sys.stderr)

def summarize(now, before, id_list:list[str], checked:str):
    newest=(now[0],now[1])
    earlier=(before[0],before[1]) if before else None
    current=cohort_ranks(now[3])
    prior=cohort_ranks(before[3]) if before else {}
    current_bw=cohort_ranks([r for r in now[3] if r.get("association")==BW_ASSOCIATION])
    prior_bw=cohort_ranks([r for r in before[3] if r.get("association")==BW_ASSOCIATION]) if before else {}
    players={}
    for pid in id_list:
        disciplines={}
        for dis in VALID_DISCIPLINES:
            r=current.get((pid,dis))
            if not r:continue
            p=prior.get((pid,dis))
            comparable=bool(p and p["ageClass"]==r["ageClass"])
            bw=current_bw.get((pid,dis))
            before_bw=prior_bw.get((pid,dis))
            comparable_bw=bool(bw and before_bw and before_bw["ageClass"]==bw["ageClass"])
            disciplines[dis]={
                "ageClass":r["ageClass"],"birthYear":r["birthYear"],
                "points":r["points"],"ageClassRank":r["ageClassRank"],
                "bwAgeClassRank":bw["ageClassRank"] if bw else None,
                "bwAgeClassSize":bw["ageClassSize"] if bw else None,
                "previousBwAgeClassRank":before_bw["ageClassRank"] if comparable_bw else None,
                "bwAgeClassChange":(before_bw["ageClassRank"]-bw["ageClassRank"]) if comparable_bw else None,
                "overallRank":r["overallRank"],"ageClassSize":r["ageClassSize"],
                "previousAgeClassRank":p["ageClassRank"] if comparable else None,
                "previousPoints":p["points"] if comparable else None,
                "ageClassChange":(p["ageClassRank"]-r["ageClassRank"]) if comparable else None,
                "tournaments":r["tournaments"]
            }
        if disciplines: players[pid]={"name":(next((r["firstName"]+" "+r["lastName"] for r in current.values() if r["id"]==pid),"")).strip(),"disciplines":disciplines}
    return {"schemaVersion":2,"status":"available","type":"dbv-published-excel-derived-ageclass-rank",
      "checkedAt":checked,
      "current":{"year":newest[0],"week":newest[1],"url":now[2]},
      "previous":{"year":earlier[0],"week":earlier[1],"url":before[2]} if before else None,
      "sourceUrl":HISTORY,"region":BW_ASSOCIATION,"players":players,
      "method":"DE: Rang in DBV-Altersklasse AKL2 (U11 umfasst mehrere Geburtsjahrgänge), Geschlecht und Disziplin; BW: gleiche Altersklasse für LVName=BAW-Baden-Württemberg. Gleichstand ergibt gleichen Rang.",
      "notes":"Offizielle veröffentlichte Excel-Ranglisten. Nur öffentliche Spieler-IDs aus der konfigurierten Familien-Auswahl werden gespeichert. Veröffentlichungszeitpunkt ist nicht notwendig Donnerstag."}

def should_publish(old:dict|None,new:dict) -> bool:
    """Do not regress a deployed ranking to an older, archived or stale snapshot."""
    if not old or old.get("status")!="available" or old.get("schemaVersion")!=2:
        return True
    previous=old.get("current") or {}
    current=new.get("current") or {}
    old_week=(int(previous.get("year") or 0),int(previous.get("week") or 0))
    new_week=(int(current.get("year") or 0),int(current.get("week") or 0))
    if new_week<old_week:
        return False
    if new_week==old_week:
        if previous.get("kind")=="current-export" and current.get("kind")!="current-export":
            return False
        if previous.get("sourceUpdatedAt") and current.get("sourceUpdatedAt") and current["sourceUpdatedAt"]<previous["sourceUpdatedAt"]:
            return False
    return True

def main():
    allow=[s.strip() for s in os.environ.get("BADMINTON_PLAYER_IDS","05-070879").split(",") if re.fullmatch(r"\d\d-\d{6}",s.strip())]
    if not allow:raise SystemExit("No valid player IDs")
    now=datetime.now(timezone.utc)
    checked=now.isoformat(timespec="seconds").replace("+00:00","Z")
    live=current_public_export()
    archives=[]
    if live:
        # For current KW41, compare to most recent *older* published weekly archive (KW40).
        for archive in published_export(now):
            if (archive[0],archive[1])<(live[0],live[1]):
                archives.append(archive)
                break
        latest=live
        previous=archives[0] if archives else None
    else:
        for archive in published_export(now):
            archives.append(archive)
            if len(archives)==2:break
        if not archives:
            print("No official Excel export was obtainable; existing snapshot left unchanged.")
            return
        latest=archives[0]
        previous=archives[1] if len(archives)>1 else None
    result=summarize(latest,previous,allow,checked)
    result["current"]["kind"]="current-export" if live else "weekly-archive"
    if live:
        result["current"]["sourceUpdatedAt"]=live[4]
        result["current"]["downloadUrl"]=live[2]
        result["current"]["url"]=OFFICIAL+"/ranking"
        result["sourceUrl"]=OFFICIAL+"/ranking"
    # Independent audit against raw official Excel rows, protecting year/gender/disc filters.
    latest=latest[3]
    for pid,player in result["players"].items():
        for disc,rank in player["disciplines"].items():
            own=next((row for row in latest if row["id"]==pid and row["discipline"]==disc),None)
            if own is None:raise ValueError("Selected player vanished from ranking input")
            higher={row["id"] for row in latest
                    if row["discipline"]==disc and row["gender"]==own["gender"]
                    and row["ageClass"]==own["ageClass"] and row["points"]>own["points"]}
            audit=1+len(higher)
            if audit!=rank["ageClassRank"]:
                raise ValueError(f"Cohort rank audit failed: {pid} {disc} got {rank['ageClassRank']}, expected {audit}")
            if own["association"]==BW_ASSOCIATION:
                higher_bw={row["id"] for row in latest
                    if row["discipline"]==disc and row["gender"]==own["gender"]
                    and row["ageClass"]==own["ageClass"]
                    and row.get("association")==BW_ASSOCIATION and row["points"]>own["points"]}
                audit_bw=1+len(higher_bw)
                if rank["bwAgeClassRank"]!=audit_bw:
                    raise ValueError(f"BW cohort audit failed: {pid} {disc} got {rank['bwAgeClassRank']}, expected {audit_bw}")
            print(f"VERIFIED {pid} {disc}: BW={rank['bwAgeClassRank']} DE={rank['ageClassRank']} ageClass={own['ageClass']} birthyear={own['birthYear']} points={own['points']} previous BW={rank['previousBwAgeClassRank']}")
    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    old=json.loads(OUTPUT.read_text()) if OUTPUT.exists() else None
    if not should_publish(old,result):
        print("Refusing to replace newer DBV ranking with an older or less reliable source; existing snapshot retained.")
        return
    if old and old.get("schemaVersion")==2 and old.get("current")==result.get("current") and old.get("previous")==result.get("previous") and old.get("players")==result.get("players"):
        print("Same official calendar-week exports; no change in KPI snapshot.")
        return
    OUTPUT.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n")
    print("Saved official weekly ranking cohort snapshot:",result["current"],"previous",result["previous"],"players:",list(result["players"]))

if __name__=="__main__":
    main()
